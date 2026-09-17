# -*- coding: utf-8 -*-
"""
憨憨保种区明细导出器 (hhclubdetail)
====================================
每天定时（默认 23:55）抓取憨憨站「个人页面 - 完成的保种区种子」全部明细，
导出为 Excel（与站点页面逐列一致），供保种区公式验证与数据积累。

为什么是 23:55：站点每日 0 点结算，0 点后再导出看到的是「次日」状态；
23:55 抓取的是当天最后状态，正好对应当日 0 点结算的输入。

输出文件（保存到插件数据目录）：
    Hhan保种区完整明细_YYYY-MM-DD.xlsx
    - Sheet1 保种区明细：序号/种子ID/种子名称/种子大小/初始保种人数/现在保种人数/完成时间/上次结算时间/今日做种时间/今日是否达标
    - Sheet2 分档统计：站点「初始保种数1人/2-3人/4-5人」的憨豆/积分统计 + 汇总
"""

import html as _html
import re
import threading
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import requests
from apscheduler.triggers.cron import CronTrigger

try:
    from fastapi import APIRouter, HTTPException
    from fastapi.responses import FileResponse
    HAS_FASTAPI = True
except Exception:
    HAS_FASTAPI = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except Exception:
    HAS_OPENPYXL = False

try:
    from app.core.config import settings
    from app.log import logger
    from app.plugins import _PluginBase
except Exception:
    settings = None
    logger = None

    class _PluginBase:
        pass


# ============================================================
# 纯函数：解析（不依赖 MP 环境，可独立测试）
# ============================================================
def _row_cells(tr_html: str) -> List[str]:
    """解析一行 <tr> 的单元格文本（去标签/去实体）"""
    cells = re.findall(r'<t[dh][^>]*>(.*?)</t[dh]>', tr_html, re.S)
    return [_html.unescape(re.sub(r'<[^>]+>', '', c)).strip() for c in cells]


def _tables(html: str) -> List[str]:
    """提取页面中所有 <table> 块"""
    return re.findall(r'<table\b[^>]*>(.*?)</table>', html, re.S | re.I)


def parse_max_page(html: str) -> int:
    """从分页链接解析最大页码（0 基）。无分页链接返回 0。

    兼容 URL 参数顺序（?action=7&id=..&page=N 或 ?id=..&action=7&page=N），
    保种区种子增多后分页数会变多，以页面分页链接为准而非写死页数。
    """
    pages = [int(p) for p in re.findall(
        r'(?:userdetails\.php)?\?[^"\']*action=7[^"\']*page=(\d+)', html)]
    return max(pages) if pages else 0


def parse_action7(html: str) -> Tuple[List[List[str]], Dict[str, Any]]:
    """解析 action=7 页面：明细表 + 分档统计（憨豆表/积分表）

    Returns:
        (rows, summary):
            rows: 每行 9 列 [种子ID, 名称, 大小, 初始人数, 现在人数, 完成时间, 上次结算, 今日做种, 是否达标]
            summary: {"bean": {...}, "pt": {...}, "count": 明细行数, "pass": 达标数}
    """
    rows: List[List[str]] = []
    bean = None
    pt = None
    for tb in _tables(html):
        trs = re.findall(r'<tr\b[^>]*>(.*?)</tr>', tb, re.S | re.I)
        if not trs:
            continue
        first = _row_cells(trs[0])
        joined = "".join(first)
        if "今日是否达标" in joined and "种子ID" in joined:
            # 明细表：跳过表头行
            for tr in trs[1:]:
                cells = _row_cells(tr)
                if len(cells) >= 9:
                    rows.append(cells[:9])
        elif "初始保种数" in joined:
            # 分档统计表：行 = [数量/体积, 1人, 2-3人, 4-5人]
            stat = {"quant": ["", "", "", ""], "vol": ["", "", "", ""]}
            for tr in trs[1:]:
                cells = _row_cells(tr)
                if len(cells) < 4:
                    continue
                key = cells[0]
                if "数量" in key:
                    stat["quant"] = cells[:4]
                elif "体积" in key:
                    stat["vol"] = cells[:4]
            if bean is None:
                bean = stat
            else:
                pt = stat
    summary: Dict[str, Any] = {"count": len(rows), "pass": 0, "bean": bean, "pt": pt}
    for r in rows:
        flag = r[8] if len(r) > 8 else ""
        if flag and "未" not in flag and "达标" in flag:
            summary["pass"] += 1
    return rows, summary


def _size_to_gb(text: str) -> float:
    m = re.search(r'([\d.]+)\s*(TB|GB|MB|KB)', text or "", re.I)
    if not m:
        return 0.0
    v = float(m.group(1))
    u = m.group(2).upper()
    if u == "TB":
        return v * 1024
    if u == "GB":
        return v
    if u == "MB":
        return v / 1024
    return v / 1048576


def write_excel(path: Path, rows: List[List[str]], summary: Dict[str, Any]) -> bool:
    """生成 xlsx：Sheet1 明细 + Sheet2 分档统计。openpyxl 不可用时降级 .csv"""
    try:
        if HAS_OPENPYXL:
            _write_xlsx(path, rows, summary)
        else:
            _write_csv(path.with_suffix(".csv"), rows, summary)
        return True
    except Exception:
        try:
            _write_csv(path.with_suffix(".csv"), rows, summary)
            return True
        except Exception:
            return False


def _write_xlsx(path: Path, rows: List[List[str]], summary: Dict[str, Any]):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "保种区明细"
    headers = ["序号", "种子ID", "种子名称", "种子大小", "初始保种人数", "现在保种人数",
               "完成时间", "上次结算时间", "今日做种时间", "今日是否达标"]
    ws.append(headers)
    widths = [6, 12, 60, 12, 14, 14, 18, 18, 14, 14]
    for idx, row in enumerate(rows, 1):
        ws.append([idx] + (row[:9] if len(row) >= 9 else row))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("分档统计")
    ws2.append(["憨豆分档统计（初始保种数）"])
    ws2.append([])
    ws2.append(["", "初始保种数1人", "初始保种数2-3人", "初始保种数4-5人"])
    if summary.get("bean"):
        ws2.append(["数量"] + summary["bean"]["quant"][1:])
        ws2.append(["体积"] + summary["bean"]["vol"][1:])
    else:
        ws2.append(["数量", "-", "-", "-"])
        ws2.append(["体积", "-", "-", "-"])
    ws2.append([])
    ws2.append(["做种积分分档统计（初始保种数）"])
    ws2.append([])
    ws2.append(["", "初始保种数1人", "初始保种数2-3人", "初始保种数4-5人"])
    if summary.get("pt"):
        ws2.append(["数量"] + summary["pt"]["quant"][1:])
        ws2.append(["体积"] + summary["pt"]["vol"][1:])
    else:
        ws2.append(["数量", "-", "-", "-"])
        ws2.append(["体积", "-", "-", "-"])
    ws2.append([])
    ws2.append(["明细总行数", summary.get("count", 0)])
    ws2.append(["今日达标数", summary.get("pass", 0)])
    for i, w in enumerate([12, 18, 18, 18], 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    wb.save(path)


def _write_csv(path: Path, rows: List[List[str]], summary: Dict[str, Any]):
    import csv
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["序号", "种子ID", "种子名称", "种子大小", "初始保种人数", "现在保种人数",
                    "完成时间", "上次结算时间", "今日做种时间", "今日是否达标"])
        for idx, row in enumerate(rows, 1):
            w.writerow([idx] + (row[:9] if len(row) >= 9 else row))
        w.writerow([])
        w.writerow(["明细总行数", summary.get("count", 0)])
        w.writerow(["今日达标数", summary.get("pass", 0)])


# ============================================================
# MP 插件
# ============================================================
class HHClubDetailExporter(_PluginBase):
    """憨憨保种区明细导出器"""

    plugin_name = "憨憨保种区明细导出"
    plugin_desc = "每天定时（默认23:55）导出憨憨站保种区保种明细为Excel，用于公式验证与数据积累。"
    plugin_icon = "https://raw.githubusercontent.com/SixOrg/MoviePilot-Plugins/main/plugins.v2/hhclubdetail/icon.png"
    plugin_version = "1.0.2"
    plugin_author = "六个橙子"
    author_url = "https://github.com/SixOrg"
    plugin_config_prefix = "hhclubdetail_"
    plugin_order = 7
    auth_level = 1

    _enabled = False
    _cron = "55 23 * * *"
    _retain_days = 30
    _notify = False
    _cookie = ""
    _site_url = "https://hhanclub.net"
    _uid: Optional[str] = None
    _last_result = "尚未运行"
    _thread: Optional[threading.Thread] = None

    def init_plugin(self, config: dict = None):
        cfg = config or {}
        self._enabled = cfg.get("enabled", False)
        self._cron = cfg.get("cron") or "55 23 * * *"
        try:
            self._retain_days = int(cfg.get("retain_days") or 30)
        except Exception:
            self._retain_days = 30
        self._notify = cfg.get("notify", False)
        self._cookie = cfg.get("cookie") or ""
        self._uid = cfg.get("uid") or None
        self._last_result = "尚未运行"
        if self._enabled:
            self._cleanup_old_files()
            # 立即运行一次
            if cfg.get("onlyonce"):
                self._thread = threading.Thread(target=self._run_worker, args=())
                self._thread.daemon = True
                self._thread.start()

    def get_state(self) -> bool:
        return self._enabled

    @staticmethod
    def get_command() -> List[Dict[str, Any]]:
        return [{
            "cmd": "/hhclubdetail_export",
            "event": None,
            "desc": "立即导出憨憨保种区明细",
            "category": "插件命令",
            "data": {}
        }]

    def get_api(self) -> List[Dict[str, Any]]:
        if not HAS_FASTAPI:
            return []
        api = APIRouter(prefix="/hhclubdetail")

        @api.post("/export")
        def api_export():
            return {"success": True, "message": self.run()}

        @api.get("/list")
        def api_list():
            files = []
            try:
                for p in sorted(self.get_data_path().glob("*.xlsx"), reverse=True)[:30]:
                    files.append({"name": p.name, "size": p.stat().st_size,
                                  "mtime": datetime.fromtimestamp(
                                      p.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")})
            except Exception as e:
                logger.error(f"列出导出文件失败：{e}")
            return {"files": files}

        @api.get("/download")
        def api_download(name: str):
            path = self.get_data_path() / name
            if not path.exists():
                raise HTTPException(status_code=404, detail="文件不存在")
            return FileResponse(path, filename=path.name)

        return [{
            "path": "/hhclubdetail",
            "endpoint": api,
            "methods": ["GET", "POST"],
            "auth": "apikey",
            "summary": "保种区明细导出API",
            "description": "POST /export 立即导出；GET /list 文件列表；GET /download?name= 下载",
        }]

    def get_service(self) -> List[Dict[str, Any]]:
        if not self.get_state():
            return []
        return [{
            "id": "HHClubDetailExporter",
            "name": "憨憨保种区明细导出",
            "trigger": CronTrigger.from_crontab(self._cron),
            "func": self.run,
            "kwargs": {}
        }]

    def get_form(self) -> Tuple[List[dict], Dict[str, Any]]:
        form = [
            {
                'component': 'VForm',
                'content': [
                    {
                        'component': 'VRow',
                        'content': [
                            {
                                'component': 'VCol',
                                'props': {'cols': 12, 'md': 4},
                                'content': [
                                    {'component': 'VSwitch', 'props': {'model': 'enabled', 'label': '启用插件'}}
                                ]
                            },
                            {
                                'component': 'VCol',
                                'props': {'cols': 12, 'md': 4},
                                'content': [
                                    {'component': 'VSwitch', 'props': {'model': 'onlyonce', 'label': '立即运行一次'}}
                                ]
                            },
                            {
                                'component': 'VCol',
                                'props': {'cols': 12, 'md': 4},
                                'content': [
                                    {'component': 'VSwitch', 'props': {'model': 'notify', 'label': '发送通知'}}
                                ]
                            }
                        ]
                    },
                    {
                        'component': 'VRow',
                        'content': [
                            {
                                'component': 'VCol',
                                'props': {'cols': 12, 'md': 6},
                                'content': [
                                    {
                                        'component': 'VCronField',
                                        'props': {
                                            'model': 'cron',
                                            'label': '导出时间',
                                            'placeholder': '55 23 * * *（每天23:55，0点结算前抓当天明细）'
                                        }
                                    }
                                ]
                            },
                            {
                                'component': 'VCol',
                                'props': {'cols': 12, 'md': 6},
                                'content': [
                                    {
                                        'component': 'VInputNumber',
                                        'props': {
                                            'model': 'retain_days',
                                            'label': '保留天数',
                                            'min': 1,
                                            'max': 365
                                        }
                                    }
                                ]
                            }
                        ]
                    },
                    {
                        'component': 'VRow',
                        'content': [
                            {
                                'component': 'VCol',
                                'props': {'cols': 12, 'md': 6},
                                'content': [
                                    {
                                        'component': 'VTextField',
                                        'props': {
                                            'model': 'uid',
                                            'label': '站点UID（可选）',
                                            'placeholder': '14332（留空则访问站点主页自动获取）'
                                        }
                                    }
                                ]
                            },
                            {
                                'component': 'VCol',
                                'props': {'cols': 12, 'md': 6},
                                'content': [
                                    {
                                        'component': 'VTextField',
                                        'props': {
                                            'model': 'cookie',
                                            'label': 'Cookie（可选）',
                                            'placeholder': '留空则自动从MP站点管理获取憨憨站Cookie'
                                        }
                                    }
                                ]
                            }
                        ]
                    },
                    {
                        'component': 'VAlert',
                        'props': {'type': 'info', 'variant': 'tonal',
                                  'text': '导出文件保存到插件数据目录，页面底部可见文件列表与完整路径。'
                                          '攒几天后把 xlsx 发给助手做公式验证。'}
                    }
                ]
            }
        ]
        return form, {
            "enabled": self._enabled,
            "cron": self._cron,
            "retain_days": self._retain_days,
            "notify": self._notify,
            "uid": self._uid or "",
            "cookie": self._cookie or "",
        }

    def get_api(self) -> List[Dict[str, Any]]:
        return [
            {
                "path": "/delete_file",
                "endpoint": self._delete_file,
                "methods": ["GET"],
                "summary": "删除指定导出文件",
                "description": "按文件名删除插件数据目录下的导出文件（仅限 .xlsx）",
            }
        ]

    def _delete_file(self, name: str = "") -> Dict[str, Any]:
        """删除指定导出文件（防路径穿越：仅允许数据目录下的 .xlsx）"""
        try:
            if not name or not name.strip():
                return {"success": False, "message": "缺少文件名"}
            data_path = self.get_data_path().resolve()
            target = (data_path / name.strip()).resolve()
            if target.parent != data_path or target.suffix.lower() != ".xlsx":
                return {"success": False, "message": "文件名无效"}
            if target.exists():
                target.unlink()
                logger.info(f"hhclubdetail - 已手动删除导出文件：{name}")
                return {"success": True, "message": f"已删除 {name}"}
            return {"success": False, "message": "文件不存在"}
        except Exception as e:
            logger.error(f"hhclubdetail - 删除文件失败：{e}")
            return {"success": False, "message": str(e)}

    def get_page(self) -> List[dict]:
        files = []
        try:
            files = sorted(self.get_data_path().glob("*.xlsx"), reverse=True)[:20]
        except Exception:
            pass
        rows = []
        if files:
            for p in files:
                mtime = datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
                size_kb = p.stat().st_size / 1024
                rows.append({
                    "component": "tr",
                    "content": [
                        {"component": "td", "text": p.name},
                        {"component": "td", "text": f"{size_kb:.0f} KB"},
                        {"component": "td", "text": mtime},
                        {
                            "component": "td",
                            "content": [{
                                "component": "VBtn",
                                "props": {"color": "error", "variant": "outlined", "size": "x-small"},
                                "text": "删除",
                                "events": {
                                    "click": {
                                        "api": "plugin/HHClubDetailExporter/delete_file",
                                        "method": "get",
                                        "params": {"name": p.name},
                                    }
                                },
                            }]
                        },
                    ],
                })
        else:
            rows.append({
                "component": "tr",
                "content": [{
                    "component": "td",
                    "text": f"暂无导出文件（每天 {self._cron} 自动导出，或勾选「立即运行一次」手动导出）",
                    "props": {"colspan": 4, "class": "text-center"},
                }]
            })
        return [
            {
                "component": "div",
                "props": {"class": "pa-2"},
                "content": [
                    {
                        "component": "p",
                        "props": {"class": "text-caption mb-1"},
                        "text": f"保存目录：{self.get_data_path()}",
                    },
                    {
                        "component": "p",
                        "props": {"class": "text-caption mb-1"},
                        "text": f"定时：{self._cron}（23:55 = 0 点结算前的当天最后状态）  保留最近 {self._retain_days} 天",
                    },
                    {
                        "component": "p",
                        "props": {"class": "text-caption mb-1"},
                        "text": f"最近运行状态：{self._last_result}",
                    },
                ],
            },
            {
                "component": "VTable",
                "props": {"hover": True},
                "content": [
                    {
                        "component": "thead",
                        "content": [
                            {
                                "component": "tr",
                                "content": [
                                    {"component": "th", "text": "文件名"},
                                    {"component": "th", "text": "大小"},
                                    {"component": "th", "text": "导出时间"},
                                    {"component": "th", "text": "操作"},
                                ],
                            }
                        ],
                    },
                    {"component": "tbody", "content": rows},
                ],
            },
        ]

    def stop_service(self):
        self._enabled = False

    # ============================================================
    # 运行逻辑
    # ============================================================
    def run(self) -> str:
        """MP 运行按钮 / cron / 命令统一入口（同步执行，返回结果摘要）"""
        self._last_result = self._export_once()
        return self._last_result

    def _run_worker(self):
        """onlyonce 后台线程入口"""
        self.run()

    def _export_once(self) -> str:
        try:
            uid = self._get_site_uid()
            if not uid:
                return "未获取到站点UID（检查Cookie/UA是否有效），导出中止"
            rows_all: List[List[str]] = []
            summary_all: Optional[Dict[str, Any]] = None
            session = self._session()
            site_url = self._get_site_url()
            # 先抓第一页，从分页链接确定总页数（不写死页数，种子增多自动跟随）
            first_url = f"{site_url}/userdetails.php?id={uid}&action=7"
            r = session.get(first_url, timeout=30)
            r.raise_for_status()
            rows, summary = parse_action7(r.text)
            if not rows:
                return "未解析到任何明细（检查Cookie是否有效/页面结构是否变化）"
            rows_all.extend(rows)
            summary_all = summary
            max_page = parse_max_page(r.text)
            page = 1
            # 安全上限 100 页（约 5000 颗），防分页解析异常导致死循环
            while page <= max_page and page < 100:
                url = f"{first_url}&page={page}"
                r = session.get(url, timeout=30)
                r.raise_for_status()
                rows, _ = parse_action7(r.text)
                if not rows:
                    break
                rows_all.extend(rows)
                page += 1
            summary_all["count"] = len(rows_all)
            summary_all["pass"] = sum(1 for row in rows_all if row[8] and "未" not in row[8] and "达标" in row[8])
            # 写文件
            data_dir = self.get_data_path()
            try:
                data_dir.mkdir(parents=True, exist_ok=True)
            except Exception:
                pass
            fname = f"Hhan保种区完整明细_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
            path = data_dir / fname
            if not write_excel(path, rows_all, summary_all):
                return "Excel 写入失败（检查磁盘权限/目录可写）"
            self._cleanup_old_files()
            msg = f"导出成功：{fname}（明细 {summary_all['count']} 行 / 达标 {summary_all['pass']} 行）"
            if self._notify:
                try:
                    self.post_message("憨憨保种区明细导出", msg)
                except Exception as e:
                    logger.error(f"发送通知失败：{e}")
            logger.info(f"hhclubdetail - {msg}")
            return msg
        except Exception as e:
            logger.error(f"憨憨保种区明细导出失败：{e}")
            return f"导出失败：{e}"

    def _cleanup_old_files(self):
        """删除超过保留天数的旧 xlsx"""
        try:
            if self._retain_days <= 0:
                return
            data_dir = self.get_data_path()
            if not data_dir.exists():
                return
            cutoff = datetime.now().timestamp() - self._retain_days * 86400
            for p in data_dir.glob("*.xlsx"):
                try:
                    if p.stat().st_mtime < cutoff:
                        p.unlink()
                        logger.info(f"hhclubdetail - 已清理旧文件：{p.name}")
                except Exception:
                    pass
        except Exception as e:
            logger.error(f"清理旧文件失败：{e}")

    # ============================================================
    # 站点抓取（复用 MP 站点管理）
    # ============================================================
    def _get_mp_site(self):
        try:
            from app.db.site_oper import SiteOper
            site_oper = SiteOper()
            for site in site_oper.list_active():
                url = site.url or ""
                domain = site.domain or ""
                if "hhanclub" in url.lower() or "hhclub" in url.lower() \
                        or "hhanclub" in domain.lower() or "hhclub" in domain.lower():
                    return site
        except Exception as e:
            logger.error(f"读取MP站点配置失败：{e}")
        return None

    def _get_site_cookie(self) -> str:
        if self._cookie:
            return self._cookie
        site = self._get_mp_site()
        if site and site.cookie:
            return site.cookie
        return ""

    def _get_site_url(self) -> str:
        site = self._get_mp_site()
        if site and site.url:
            return site.url.rstrip("/")
        return self._site_url

    def _get_site_uid(self) -> Optional[str]:
        """获取站点UID：优先配置缓存值，其次访问站点主页自动解析"""
        if self._uid:
            return self._uid
        try:
            session = self._session()
            url = self._get_site_url()
            r = session.get(url, timeout=30)
            m = re.search(r"userdetails\.php\?id=(\d+)", r.text)
            if m:
                self._uid = m.group(1)
                logger.info(f"hhclubdetail - 已从站点主页自动获取UID：{self._uid}")
                return self._uid
        except Exception as e:
            logger.error(f"自动获取站点UID失败：{e}（如站点主页访问超时，请在插件配置中直接填写UID，如14332）")
        return None

    def _session(self) -> requests.Session:
        s = requests.Session()
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                          "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        }
        site = self._get_mp_site()
        if site and site.ua:
            headers["User-Agent"] = site.ua
        cookie = self._get_site_cookie()
        if cookie:
            headers["Cookie"] = cookie
        s.headers.update(headers)
        return s
